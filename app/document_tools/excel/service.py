from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

from app.document_tools.common import MIME_TYPES, ensure_suffix, store_document_result
from app.document_tools.schemas import ConversionResult, DocumentConversionError
from app.document_tools.storage import DocumentStorage


class ExcelToolService:
    def __init__(self, storage: DocumentStorage):
        self.storage = storage

    def create(
        self,
        *,
        user_id: str,
        filename: str = "",
        sheets: list[dict[str, Any]] | None = None,
    ) -> ConversionResult:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise DocumentConversionError("服务器未安装 openpyxl，暂时无法生成 Excel。", status_code=503) from exc

        workbook = Workbook()
        sheet_specs = sheets or [{"name": "Sheet1", "rows": []}]
        for index, spec in enumerate(sheet_specs):
            worksheet = workbook.active if index == 0 else workbook.create_sheet()
            worksheet.title = str(spec.get("name") or f"Sheet{index + 1}")[:31]
            rows = spec.get("rows") or []
            for row_index, row in enumerate(rows, start=1):
                for column_index, value in enumerate(row, start=1):
                    cell = worksheet.cell(row=row_index, column=column_index, value=value)
                    if row_index == 1:
                        cell.font = Font(bold=True)
                        cell.fill = PatternFill("solid", fgColor="D9EAF7")
                        cell.alignment = Alignment(horizontal="center")
            for column_index in range(1, worksheet.max_column + 1):
                values = [worksheet.cell(row=row, column=column_index).value for row in range(1, worksheet.max_row + 1)]
                width = max([len(str(value)) for value in values if value is not None] or [8])
                worksheet.column_dimensions[get_column_letter(column_index)].width = min(width + 3, 60)
        return self._save_workbook(user_id, workbook, filename, "excel_create", "已创建 Excel 工作簿。")

    def read(self, *, user_id: str, file_id: str) -> dict[str, Any]:
        stored = self.storage.get_file(user_id, file_id)
        if stored is None:
            raise DocumentConversionError("找不到 Excel 文件，请重新上传或生成后再试。", status_code=404)
        from openpyxl import load_workbook

        workbook = load_workbook(stored.path, data_only=False)
        try:
            sheets = {}
            for sheet_name in workbook.sheetnames:
                worksheet = workbook[sheet_name]
                rows = []
                for row in worksheet.iter_rows(values_only=True):
                    rows.append(["" if value is None else value for value in row])
                sheets[sheet_name] = rows
            return {"file_id": stored.file_id, "filename": stored.filename, "sheets": sheets}
        finally:
            workbook.close()

    def calculate(
        self,
        *,
        user_id: str,
        file_id: str,
        operations: list[dict[str, Any]],
        filename: str = "",
    ) -> ConversionResult:
        stored = self.storage.get_file(user_id, file_id)
        if stored is None:
            raise DocumentConversionError("找不到 Excel 文件，请重新上传或生成后再试。", status_code=404)
        from openpyxl import load_workbook

        workbook = load_workbook(stored.path)
        try:
            for operation in operations:
                sheet = workbook[str(operation.get("sheet") or workbook.sheetnames[0])]
                op_type = str(operation.get("type") or "")
                if op_type == "formula":
                    cell = str(operation.get("cell") or "")
                    if not cell:
                        continue
                    if "value" in operation:
                        sheet[cell] = operation.get("value")
                    else:
                        formula = str(operation.get("formula") or "")
                        sheet[cell] = formula if formula.startswith("=") else f"={formula}"
                elif op_type == "sum":
                    target_cell = str(operation.get("target_cell") or "")
                    range_ref = str(operation.get("range") or "")
                    if target_cell and range_ref:
                        sheet[target_cell] = f"=SUM({range_ref})"
            return self._save_workbook(
                user_id,
                workbook,
                filename or f"{Path(stored.filename).stem}-calculated.xlsx",
                "excel_calculate",
                "已完成 Excel 运算设置。",
            )
        finally:
            workbook.close()

    def format_table(
        self,
        *,
        user_id: str,
        file_id: str,
        sheet: str,
        range_ref: str,
        filename: str = "",
    ) -> ConversionResult:
        stored = self.storage.get_file(user_id, file_id)
        if stored is None:
            raise DocumentConversionError("找不到 Excel 文件，请重新上传或生成后再试。", status_code=404)
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo

        workbook = load_workbook(stored.path)
        try:
            worksheet = workbook[sheet]
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="4472C4")
                cell.alignment = Alignment(horizontal="center")
            table = Table(displayName=f"Table_{worksheet.max_row}_{worksheet.max_column}", ref=range_ref)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium9",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            if not worksheet.tables:
                worksheet.add_table(table)
            return self._save_workbook(
                user_id,
                workbook,
                filename or f"{Path(stored.filename).stem}-formatted.xlsx",
                "excel_format_table",
                "已格式化 Excel 表格。",
            )
        finally:
            workbook.close()

    def _save_workbook(self, user_id: str, workbook: Any, filename: str, operation: str, message: str) -> ConversionResult:
        buffer = BytesIO()
        workbook.save(buffer)
        return store_document_result(
            self.storage,
            user_id=user_id,
            filename=ensure_suffix(filename, ".xlsx", "workbook"),
            mime_type=MIME_TYPES["xlsx"],
            content=buffer.getvalue(),
            operation=operation,
            message=message,
        )
