from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

from app.document_tools.common import MIME_TYPES, ensure_suffix, store_document_result
from app.document_tools.schemas import ConversionResult, DocumentConversionError
from app.document_tools.storage import DocumentStorage


def _set_run_font(run: Any, font_name: str | None, size_pt: float | None = None) -> None:
    if font_name:
        run.font.name = font_name
        try:
            from docx.oxml.ns import qn

            run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        except Exception:
            pass
    if size_pt:
        from docx.shared import Pt

        run.font.size = Pt(size_pt)


def _apply_paragraph_style(paragraph: Any, style: dict[str, Any]) -> None:
    alignment = str(style.get("alignment") or "").lower()
    if alignment:
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        alignments = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
            "左": WD_ALIGN_PARAGRAPH.LEFT,
            "居中": WD_ALIGN_PARAGRAPH.CENTER,
            "右": WD_ALIGN_PARAGRAPH.RIGHT,
            "两端": WD_ALIGN_PARAGRAPH.JUSTIFY,
        }
        for key, value in alignments.items():
            if key in alignment:
                paragraph.alignment = value
                break
    line_spacing = style.get("line_spacing")
    if line_spacing:
        paragraph.paragraph_format.line_spacing = float(line_spacing)


class WordToolService:
    def __init__(self, storage: DocumentStorage):
        self.storage = storage

    def create(
        self,
        *,
        user_id: str,
        filename: str = "",
        title: str = "",
        blocks: list[dict[str, Any]] | None = None,
        style: dict[str, Any] | None = None,
    ) -> ConversionResult:
        try:
            from docx import Document
        except ImportError as exc:
            raise DocumentConversionError("服务器未安装 python-docx，暂时无法生成 Word。", status_code=503) from exc

        doc = Document()
        style = style or {}
        body_font = style.get("body_font") or ("宋体" if style.get("preset") == "chinese_report" else None)
        body_size_pt = float(style.get("body_size_pt") or 12)
        if body_font:
            normal = doc.styles["Normal"]
            normal.font.name = body_font
            try:
                from docx.oxml.ns import qn

                normal._element.rPr.rFonts.set(qn("w:eastAsia"), body_font)
            except Exception:
                pass
        normal = doc.styles["Normal"]
        from docx.shared import Pt

        normal.font.size = Pt(body_size_pt)
        if title:
            heading = doc.add_heading(title, level=0)
            for run in heading.runs:
                _set_run_font(run, style.get("title_font") or "黑体", float(style.get("title_size_pt") or 16))

        for block in blocks or []:
            block_type = str(block.get("type") or "paragraph")
            if block_type == "heading":
                paragraph = doc.add_heading(str(block.get("text") or ""), level=int(block.get("level") or 1))
                for run in paragraph.runs:
                    _set_run_font(run, style.get("heading_font") or "黑体", None)
            elif block_type == "table":
                headers = [str(value) for value in block.get("headers") or []]
                rows = block.get("rows") or []
                table_rows = [headers] + rows if headers else rows
                if not table_rows:
                    continue
                column_count = max(len(row) for row in table_rows)
                table = doc.add_table(rows=len(table_rows), cols=column_count)
                table.style = str(block.get("style") or "Table Grid")
                for row_index, row in enumerate(table_rows):
                    for column_index, value in enumerate(row):
                        cell = table.cell(row_index, column_index)
                        cell.text = "" if value is None else str(value)
                        if row_index == 0 and headers:
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    run.bold = True
            else:
                paragraph = doc.add_paragraph(str(block.get("text") or ""))
                _apply_paragraph_style(paragraph, style)
                for run in paragraph.runs:
                    _set_run_font(run, body_font, body_size_pt)

        buffer = BytesIO()
        doc.save(buffer)
        return store_document_result(
            self.storage,
            user_id=user_id,
            filename=ensure_suffix(filename, ".docx", "document"),
            mime_type=MIME_TYPES["docx"],
            content=buffer.getvalue(),
            operation="word_create",
            message="已创建 Word 文档。",
        )

    def read(self, *, user_id: str, file_id: str) -> dict[str, Any]:
        stored = self.storage.get_file(user_id, file_id)
        if stored is None:
            raise DocumentConversionError("找不到 Word 文件，请重新上传或生成后再试。", status_code=404)
        try:
            from docx import Document
        except ImportError as exc:
            raise DocumentConversionError("服务器未安装 python-docx，暂时无法读取 Word。", status_code=503) from exc
        doc = Document(stored.path)
        paragraphs = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
        tables = []
        for table in doc.tables:
            tables.append([[cell.text for cell in row.cells] for row in table.rows])
        return {
            "file_id": stored.file_id,
            "filename": stored.filename,
            "text": "\n\n".join(paragraphs),
            "tables": tables,
            "paragraph_count": len(paragraphs),
        }

    def edit_text(
        self,
        *,
        user_id: str,
        file_id: str,
        replacements: list[dict[str, str]],
        filename: str = "",
    ) -> ConversionResult:
        stored = self.storage.get_file(user_id, file_id)
        if stored is None:
            raise DocumentConversionError("找不到 Word 文件，请重新上传或生成后再试。", status_code=404)
        from docx import Document

        doc = Document(stored.path)
        for paragraph in doc.paragraphs:
            for replacement in replacements:
                old = replacement.get("old", "")
                new = replacement.get("new", "")
                if old and old in paragraph.text:
                    paragraph.text = paragraph.text.replace(old, new)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for replacement in replacements:
                        old = replacement.get("old", "")
                        new = replacement.get("new", "")
                        if old and old in cell.text:
                            cell.text = cell.text.replace(old, new)
        return self._save_doc(user_id, doc, filename or f"{Path(stored.filename).stem}-edited.docx", "word_edit_text", "已编辑 Word 文本。")

    def format(
        self,
        *,
        user_id: str,
        file_id: str,
        style: dict[str, Any],
        filename: str = "",
    ) -> ConversionResult:
        stored = self.storage.get_file(user_id, file_id)
        if stored is None:
            raise DocumentConversionError("找不到 Word 文件，请重新上传或生成后再试。", status_code=404)
        from docx import Document

        doc = Document(stored.path)
        body_font = style.get("body_font")
        body_size_pt = float(style.get("body_size_pt") or 0) or None
        for paragraph in doc.paragraphs:
            _apply_paragraph_style(paragraph, style)
            for run in paragraph.runs:
                _set_run_font(run, body_font, body_size_pt)
        return self._save_doc(user_id, doc, filename or f"{Path(stored.filename).stem}-formatted.docx", "word_format", "已调整 Word 格式。")

    def _save_doc(self, user_id: str, doc: Any, filename: str, operation: str, message: str) -> ConversionResult:
        buffer = BytesIO()
        doc.save(buffer)
        return store_document_result(
            self.storage,
            user_id=user_id,
            filename=ensure_suffix(filename, ".docx", "document"),
            mime_type=MIME_TYPES["docx"],
            content=buffer.getvalue(),
            operation=operation,
            message=message,
        )
