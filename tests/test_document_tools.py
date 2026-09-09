from io import BytesIO
import json
from pathlib import Path

from fastapi import FastAPI
from fastapi.testclient import TestClient
from PIL import Image
from openpyxl import load_workbook

from app.document_tools.api.document_router import router
from app.document_tools.service import DocumentConversionService, infer_conversion_plan
from app.document_tools.storage import DocumentStorage
from app.main import app as main_app


ROOT = Path(__file__).resolve().parents[1]


def make_png_bytes() -> bytes:
    image = Image.new("RGB", (24, 24), "white")
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def test_document_storage_saves_upload_with_metadata(tmp_path):
    storage = DocumentStorage(tmp_path)

    stored = storage.save_upload(
        user_id="user-1",
        filename="scan page.png",
        mime_type="image/png",
        content=make_png_bytes(),
    )

    assert stored.file_id.startswith("file-")
    assert stored.filename == "scan page.png"
    assert stored.mime_type == "image/png"
    assert stored.size_bytes > 0
    assert storage.get_file("user-1", stored.file_id) == stored
    assert storage.get_file("user-2", stored.file_id) is None


def test_document_storage_rejects_metadata_paths_outside_user_directory(tmp_path):
    storage = DocumentStorage(tmp_path)
    stored = storage.save_upload(
        user_id="user-1",
        filename="scan.png",
        mime_type="image/png",
        content=make_png_bytes(),
    )
    outside_path = tmp_path / "outside.txt"
    outside_path.write_text("secret", encoding="utf-8")
    metadata_path = stored.path.parent / "metadata.json"
    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    metadata["path"] = str(outside_path)
    metadata_path.write_text(json.dumps(metadata), encoding="utf-8")

    assert storage.get_file("user-1", stored.file_id) is None


def test_infer_conversion_plan_for_core_formats():
    assert infer_conversion_plan(["photo.png"], "转成 PDF").operation == "images_to_pdf"
    assert infer_conversion_plan(["report.docx"], "转 PDF").operation == "office_to_pdf"
    assert infer_conversion_plan(["scan.pdf"], "扫描识别为 PDF").operation == "ocr_pdf"
    assert infer_conversion_plan(["book.pdf"], "转 Word").operation == "pdf_to_docx"


def test_document_conversion_service_converts_image_to_pdf(tmp_path):
    storage = DocumentStorage(tmp_path)
    uploaded = storage.save_upload(
        user_id="user-1",
        filename="receipt.png",
        mime_type="image/png",
        content=make_png_bytes(),
    )
    service = DocumentConversionService(storage)

    result = service.convert(
        user_id="user-1",
        file_ids=[uploaded.file_id],
        instruction="转成 PDF",
    )

    output = storage.get_file("user-1", result.file_id)
    assert output is not None
    assert output.filename == "receipt.pdf"
    assert output.mime_type == "application/pdf"
    assert output.path.read_bytes().startswith(b"%PDF")
    assert result.attachment == {
        "platform": "document",
        "media_type": "file",
        "title": "receipt.pdf",
        "author": "",
        "cover": "",
        "video_url": "",
        "source_url": "",
        "images": [],
        "source_images": [],
        "parse_id": result.file_id,
        "file_id": result.file_id,
        "filename": "receipt.pdf",
        "mime_type": "application/pdf",
        "download_url": f"/documents/files/{result.file_id}/download",
        "preview_url": f"/documents/files/{result.file_id}/download",
        "size_bytes": output.size_bytes,
        "status": "ready",
        "message": "已转换为 PDF。",
        "operation": "images_to_pdf",
    }


def test_document_service_creates_and_reads_word_document(tmp_path):
    storage = DocumentStorage(tmp_path)
    service = DocumentConversionService(storage)

    result = service.create_word(
        user_id="user-1",
        filename="sales-report.docx",
        title="Sales Report",
        blocks=[
            {"type": "heading", "text": "Summary", "level": 1},
            {"type": "paragraph", "text": "Revenue increased this month."},
            {
                "type": "table",
                "headers": ["Month", "Revenue"],
                "rows": [["August", 10000], ["September", 11200]],
            },
        ],
        style={"preset": "chinese_report", "body_size_pt": 12},
    )

    stored = storage.get_file("user-1", result.file_id)
    assert stored is not None
    assert stored.filename == "sales-report.docx"
    assert stored.mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert result.operation == "word_create"
    assert result.attachment["download_url"] == f"/documents/files/{result.file_id}/download"

    read_result = service.read_word(user_id="user-1", file_id=result.file_id)
    assert "Sales Report" in read_result["text"]
    assert "Revenue increased this month." in read_result["text"]
    assert read_result["tables"] == [[["Month", "Revenue"], ["August", "10000"], ["September", "11200"]]]


def test_document_service_edits_word_text_and_applies_format(tmp_path):
    storage = DocumentStorage(tmp_path)
    service = DocumentConversionService(storage)
    created = service.create_word(
        user_id="user-1",
        filename="draft.docx",
        title="Draft",
        blocks=[{"type": "paragraph", "text": "Old customer name"}],
    )

    edited = service.edit_word_text(
        user_id="user-1",
        file_id=created.file_id,
        replacements=[{"old": "Old customer name", "new": "New customer name"}],
        filename="draft-edited.docx",
    )
    formatted = service.format_word(
        user_id="user-1",
        file_id=edited.file_id,
        style={"body_font": "Calibri", "body_size_pt": 14, "line_spacing": 1.5},
        filename="draft-formatted.docx",
    )

    read_result = service.read_word(user_id="user-1", file_id=formatted.file_id)
    assert "New customer name" in read_result["text"]
    assert "Old customer name" not in read_result["text"]
    assert formatted.operation == "word_format"


def test_document_service_creates_reads_calculates_and_formats_excel(tmp_path):
    storage = DocumentStorage(tmp_path)
    service = DocumentConversionService(storage)

    created = service.create_excel(
        user_id="user-1",
        filename="sales.xlsx",
        sheets=[
            {
                "name": "Sales",
                "rows": [
                    ["Month", "Revenue", "Cost"],
                    ["August", 10000, 7000],
                    ["September", 11200, 7600],
                ],
            }
        ],
    )
    calculated = service.calculate_excel(
        user_id="user-1",
        file_id=created.file_id,
        operations=[
            {"type": "formula", "sheet": "Sales", "cell": "D1", "value": "Profit"},
            {"type": "formula", "sheet": "Sales", "cell": "D2", "formula": "B2-C2"},
            {"type": "formula", "sheet": "Sales", "cell": "D3", "formula": "B3-C3"},
            {"type": "sum", "sheet": "Sales", "range": "B2:B3", "target_cell": "B4"},
        ],
        filename="sales-calculated.xlsx",
    )
    formatted = service.format_excel_table(
        user_id="user-1",
        file_id=calculated.file_id,
        sheet="Sales",
        range_ref="A1:D4",
        filename="sales-formatted.xlsx",
    )

    stored = storage.get_file("user-1", formatted.file_id)
    assert stored is not None
    workbook = load_workbook(stored.path, data_only=False)
    sheet = workbook["Sales"]
    assert sheet["D2"].value == "=B2-C2"
    assert sheet["B4"].value == "=SUM(B2:B3)"
    assert sheet.tables
    workbook.close()

    read_result = service.read_excel(user_id="user-1", file_id=formatted.file_id)
    assert read_result["sheets"]["Sales"][0] == ["Month", "Revenue", "Cost", "Profit"]


def test_document_service_creates_and_reads_pdf(tmp_path):
    storage = DocumentStorage(tmp_path)
    service = DocumentConversionService(storage)

    result = service.create_pdf(
        user_id="user-1",
        filename="summary.pdf",
        title="Monthly Summary",
        blocks=[{"type": "paragraph", "text": "The team shipped the first document tools layer."}],
    )

    stored = storage.get_file("user-1", result.file_id)
    assert stored is not None
    assert stored.path.read_bytes().startswith(b"%PDF")
    assert result.operation == "pdf_create"

    read_result = service.read_pdf(user_id="user-1", file_id=result.file_id)
    assert "Monthly Summary" in read_result["text"]
    assert read_result["pages"] >= 1


def test_document_router_upload_convert_and_download(monkeypatch, tmp_path):
    monkeypatch.setattr(
        "app.document_tools.api.document_router.get_current_user_id",
        lambda request: "user-1",
    )
    app = FastAPI()
    app.state.document_storage = DocumentStorage(tmp_path)
    app.include_router(router)
    client = TestClient(app)

    upload_response = client.post(
        "/files/upload",
        files={"file": ("receipt.png", make_png_bytes(), "image/png")},
        headers={"Authorization": "Bearer token"},
    )

    assert upload_response.status_code == 200
    uploaded = upload_response.json()
    assert uploaded["filename"] == "receipt.png"
    assert uploaded["file_id"].startswith("file-")

    convert_response = client.post(
        "/documents/convert",
        json={"file_ids": [uploaded["file_id"]], "instruction": "转成 PDF"},
        headers={"Authorization": "Bearer token"},
    )

    assert convert_response.status_code == 200
    converted = convert_response.json()
    assert converted["attachment"]["filename"] == "receipt.pdf"
    assert converted["attachment"]["download_url"].startswith("/documents/files/")

    download_response = client.get(
        converted["attachment"]["download_url"],
        headers={"Authorization": "Bearer token"},
    )

    assert download_response.status_code == 200
    assert download_response.content.startswith(b"%PDF")


def test_document_router_rejects_oversized_upload(monkeypatch, tmp_path):
    monkeypatch.setattr(
        "app.document_tools.api.document_router.get_current_user_id",
        lambda request: "user-1",
    )
    monkeypatch.setattr(
        "app.document_tools.api.document_router.settings.document_max_upload_mb",
        1,
    )
    app = FastAPI()
    app.state.document_storage = DocumentStorage(tmp_path)
    app.include_router(router)
    client = TestClient(app)

    response = client.post(
        "/files/upload",
        files={"file": ("large.bin", b"x" * (1024 * 1024 + 1), "application/octet-stream")},
        headers={"Authorization": "Bearer token"},
    )

    assert response.status_code == 413
    assert response.json()["detail"] == "上传文件超过 1 MB 限制"


def test_main_app_mounts_document_tool_routes():
    route_paths = set(main_app.openapi()["paths"])

    assert "/files/upload" in route_paths
    assert "/documents/convert" in route_paths
    assert "/documents/files/{file_id}/download" in route_paths


def test_pyproject_declares_document_tool_dependencies():
    text = (ROOT / "pyproject.toml").read_text(encoding="utf-8")

    for dependency in (
        "python-multipart",
        "pillow",
        "python-docx",
        "pdf2docx",
        "openpyxl",
        "reportlab",
        "pypdf",
        "pdfplumber",
    ):
        assert f'"{dependency}"' in text


def test_env_example_includes_document_tool_settings():
    text = (ROOT / ".env.example").read_text(encoding="utf-8")

    for name in (
        "DOCUMENT_MAX_UPLOAD_MB",
        "DOCUMENT_CONVERSION_TIMEOUT_SECONDS",
        "DOCUMENT_OCR_TIMEOUT_SECONDS",
        "DOCUMENT_LIBREOFFICE_PATH",
        "DOCUMENT_OCRMYPDF_PATH",
        "DOCUMENT_TESSERACT_PATH",
        "DOCUMENT_OCR_LANGUAGES",
    ):
        assert name in text
